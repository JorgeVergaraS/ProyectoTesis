from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "Nexo-Cloud-Native-Informe.docx"
EVIDENCE = ROOT / "docs" / "evidence"
NAVY = "293E78"
VIOLET = RGBColor(73, 65, 180)
GRAY = RGBColor(89, 89, 89)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def cell_text(cell, value, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell_text(table.rows[0].cells[index], header, True, (255, 255, 255))
        shade(table.rows[0].cells[index], NAVY)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell_text(cells[index], value)
    return table


def heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)


def caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def add_evidence(doc, filename, text):
    path = EVIDENCE / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.35))
        picture = doc.paragraphs[-1]
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        caption(doc, text)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("NEXO")
run.bold = True
run.font.size = Pt(32)
run.font.color.rgb = VIOLET
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Integración cloud-native de autenticación y autorización")
run.bold = True
run.font.size = Pt(19)
run.font.color.rgb = RGBColor(41, 62, 120)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("DSY1107 · Informe técnico · 2 de septiembre de 2026").font.size = Pt(11)

notice = doc.add_paragraph()
notice.add_run("Estado del documento. ").bold = True
notice.add_run(
    "Registra únicamente resultados comprobados. La integración local está terminada; "
    "el acceso Microsoft real y la infraestructura AWS permanecen marcados como pendientes "
    "hasta realizar la sesión interactiva y el despliegue autorizado."
)

heading(doc, "1. Resumen ejecutivo")
doc.add_paragraph(
    "Nexo es una comunidad universitaria existente construida con Angular, Spring Boot y "
    "PostgreSQL. Se incorporó autenticación local y Microsoft Entra ID, protección de rutas, "
    "envío automático del Access Token y validación JWT en profundidad. La funcionalidad "
    "seleccionada es la consulta del perfil autenticado mediante GET /api/users/me."
)
add_table(doc, ["Componente", "Resultado comprobado", "Pendiente real"], [
    ("Angular/MSAL", "Build correcto; 30/30 pruebas", "Login real y logout en Brave"),
    ("Spring Boot/JWT", "24/24 pruebas; matriz 200/401/403", "Token real emitido por Entra"),
    ("PostgreSQL", "Usuario local persistido; Entra sin password_hash", "Upsert con cuenta real"),
    ("Azure", "IDs, URI y scope documentados", "Revalidar portal y consentimiento interactivo"),
    ("AWS", "Arquitectura y criterios definidos", "EC2, HTTP API Gateway y JWT Authorizer"),
])

heading(doc, "2. Planificación inicial")
doc.add_paragraph("Objetivo: adaptar la solución FullStack existente sin crear una aplicación nueva.")
add_table(doc, ["Responsable", "Responsabilidad", "Entregable"], [
    ("Jorge Vergara Stuardo", "Angular/MSAL, rutas, pruebas y coordinación Git", "PR frontend y evidencias UI"),
    ("Segundo integrante (por confirmar)", "Spring Security, EC2, API Gateway y pruebas API", "PR backend y evidencias cloud"),
    ("Ambos", "Revisión cruzada, informe y presentación", "Entrega y exposición de 5–10 min"),
])
doc.add_paragraph(
    "Semana 1: diagnóstico local, MSAL, Spring Security, persistencia y pruebas. "
    "Semana 2: EC2, API Gateway, CORS, JWT Authorizer, evidencias, documentación y PR."
)

heading(doc, "3. Arquitectura objetivo")
architecture = doc.add_paragraph()
architecture.alignment = WD_ALIGN_PARAGRAPH.CENTER
architecture.add_run(
    "Angular SPA\nMSAL · Authorization Code + PKCE\n"
    "↓  Authorization: Bearer <JWT redactado>\n"
    "Amazon HTTP API Gateway · CORS · JWT Authorizer\n↓\n"
    "AWS EC2 · Spring Boot Resource Server\n↓\nPostgreSQL"
).bold = True
doc.add_paragraph(
    "Microsoft Entra ID emite el token para access_as_user. API Gateway realizará la primera "
    "validación en nube y Spring Security conserva la defensa en profundidad. El despliegue AWS "
    "aún no se presenta como ejecutado."
)

heading(doc, "4. Configuración pública de Microsoft Entra ID")
add_table(doc, ["Parámetro", "Valor"], [
    ("Tenant ID", "21a4bbb2-fc48-4053-a98e-b805aa2306cc"),
    ("Frontend Client ID", "480a8cf4-c729-4ca1-8043-6c1198dfaceb"),
    ("API Client ID / audience", "0f2d7cee-cabb-4482-900d-64fb07f5f81d"),
    ("Redirect URI SPA", "http://localhost:4200"),
    ("Scope", "api://0f2d7cee-cabb-4482-900d-64fb07f5f81d/access_as_user"),
])
doc.add_paragraph(
    "La SPA no contiene Client Secret. No se usa http://127.0.0.1:4200 como redirect URI. "
    "No se modifican permisos ni se acepta consentimiento sin intervención del usuario."
)
add_evidence(doc, "01-nexo-frontend-overview.png", "Evidencia histórica disponible: overview de Nexo Frontend.")
add_evidence(doc, "02-nexo-api-scope.png", "Evidencia histórica disponible: scope access_as_user de la API Nexo.")

heading(doc, "5. Implementación Angular y MSAL")
for item in [
    "PublicClientApplication se inicializa una sola vez con tenant, client ID y redirect URI públicos.",
    "loginRedirect solicita el scope completo access_as_user.",
    "handleRedirectObservable procesa el callback y establece la cuenta activa.",
    "Las sesiones demo, local y Microsoft están diferenciadas; el callback navega a /home.",
    "acquireTokenSilent obtiene el Access Token desde la caché de MSAL.",
    "MsalAuthInterceptor agrega Bearer únicamente a la API configurada.",
    "/login es pública; /home y /profile están protegidas.",
    "La interfaz no muestra ni almacena manualmente el JWT completo.",
]:
    doc.add_paragraph(item, style="List Bullet")
add_evidence(doc, "04-nexo-login-rendered.png", "Pantalla de acceso de Nexo con alternativa local y Microsoft.")
add_evidence(doc, "05-login-validation-errors.png", "Validación de correo y contraseña mínima de 8 caracteres.")
add_evidence(doc, "06-login-success.png", "Login local exitoso y redirección a la zona protegida.")

heading(doc, "6. Spring Security y persistencia")
doc.add_paragraph(
    "El backend funciona como OAuth2 Resource Server. El JwtDecoder de Entra usa las claves "
    "públicas del tenant y valida firma, issuer, audience y timestamps. Los claims scp/scope se "
    "convierten a SCOPE_* y roles a ROLE_*. GET /api/users/me exige ROLE_USER local o "
    "SCOPE_access_as_user."
)
doc.add_paragraph(
    "Para Microsoft, oid identifica al usuario. El backend crea o actualiza su perfil y deja "
    "password_hash en NULL. No enlaza automáticamente identidades de proveedores distintos por correo."
)

heading(doc, "7. Matriz de pruebas")
add_table(doc, ["Caso", "Esperado", "Resultado"], [
    ("Registro local válido", "201", "Aprobado"),
    ("Login local válido", "200", "Aprobado"),
    ("Login local inválido", "401", "Aprobado"),
    ("Correo o contraseña inválidos", "400", "Aprobado"),
    ("GET /api/users/me sin token", "401", "Aprobado"),
    ("JWT malformado o expirado", "401", "Aprobado"),
    ("JWT válido sin permiso", "403", "Aprobado"),
    ("JWT válido con scope/rol", "200", "Aprobado en pruebas automatizadas"),
    ("Login Microsoft real", "Callback + token", "Pendiente de sesión interactiva"),
    ("API Gateway", "401 / 403 / 200", "Pendiente de despliegue AWS"),
])
doc.add_paragraph(
    "Suites: frontend 30/30; backend 24/24. El build Angular finalizó con un warning no bloqueante "
    "de presupuesto inicial (55,93 kB sobre 500 kB)."
)
doc.add_paragraph("Cabecera segura: Authorization: Bearer eyJhbGciOiJI...<redactado>.")

heading(doc, "8. Comandos reproducibles")
for command in [
    "docker compose up -d --wait",
    "cd backend; .\\mvnw.cmd verify",
    "cd frontend; corepack npm@11.13.0 ci",
    "cd frontend; corepack npm@11.13.0 run test:ci",
    "cd frontend; corepack npm@11.13.0 run build",
    "cd backend; .\\mvnw.cmd spring-boot:run",
    "cd frontend; corepack npm@11.13.0 start",
]:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(command).font.name = "Consolas"

heading(doc, "9. Riesgos, problemas y soluciones")
add_table(doc, ["Riesgo o problema", "Tratamiento"], [
    ("Bucle OAuth", "Callback centralizado, cuenta activa y navegación controlada"),
    ("Audience/issuer incorrecto", "Validadores explícitos y pruebas unitarias"),
    ("Scope ausente", "Solicitud y autorización SCOPE_access_as_user"),
    ("Mezcla de sesiones", "Proveedor demo/local/Microsoft separado"),
    ("Exposición de secretos", "Sin Client Secret; JWT y credenciales redactados"),
    ("Rúbrica menciona microservicios", "Confirmar con docente antes de dividir el monolito"),
    ("Autorregistro no disponible", "Confirmar si basta cuenta institucional precreada"),
])

doc.add_page_break()
heading(doc, "10. Git y trazabilidad")
doc.add_paragraph(
    "Frontend y documentación: feature/auth-msal. Backend: feature/backend-entra-jwt. "
    "Se preparan commits Conventional Commits y PR separados hacia develop. La publicación y "
    "solicitud de revisión se harán solo con confirmación final del usuario."
)

heading(doc, "11. Pendientes reales antes de la entrega final")
for item in [
    "Habilitar el control de Brave y completar manualmente autenticación/MFA Microsoft.",
    "Capturar callback, logout y Network con Authorization redactado.",
    "Revalidar SPA, redirect URI, permisos delegados, cliente autorizado y consentimiento.",
    "Confirmar cuenta y región AWS antes de crear recursos con costo.",
    "Desplegar Spring Boot en EC2 y configurar HTTP API Gateway, CORS, stage y JWT Authorizer.",
    "Ejecutar en Gateway la matriz 401/401/403/200 y capturar evidencias.",
    "Confirmar nombre del segundo integrante y revisor de los PR.",
]:
    doc.add_paragraph(item, style="List Bullet")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Nexo · Informe cloud-native · Sin secretos ni tokens completos").font.size = Pt(8)
doc.save(OUT)
print(OUT)
